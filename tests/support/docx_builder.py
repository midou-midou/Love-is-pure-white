"""构造测试用 DOCX 片段，用于覆盖长句、短答、重复文本与多角色对话。"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

__all__ = ["Block", "build_docx", "paragraph", "heading", "list_item", "blank"]


@dataclass(frozen=True)
class Block:
    """一个待写入的块；``text`` 中的 ``\\n`` 表示段内软换行。"""

    text: str
    kind: str = "paragraph"  # paragraph | heading | list | blank
    level: int = 0
    numbered: bool = False


def paragraph(text: str) -> Block:
    return Block(text=text)


def heading(text: str) -> Block:
    return Block(text=text, kind="heading")


def list_item(text: str, *, level: int = 0) -> Block:
    return Block(text=text, kind="list", level=level)


def blank() -> Block:
    return Block(text="", kind="blank")


def build_docx(path: Path, blocks: Sequence[Block]) -> Path:
    """写入 DOCX；列表使用真实 numPr，标题使用加粗短行（不依赖 Heading 样式）。"""
    document = Document()
    for block in blocks:
        if block.kind == "blank":
            document.add_paragraph("")
            continue
        if block.kind == "list":
            target = document.add_paragraph(style="List Number")
            _set_list_level(target, block.level)
        else:
            target = document.add_paragraph()
        _write_text(target, block.text)
        if block.kind == "heading":
            for run in target.runs:
                run.bold = True
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return Path(path)


def _write_text(target: object, text: str) -> None:
    lines = text.split("\n")
    for index, line in enumerate(lines):
        run = target.add_run(line)  # type: ignore[attr-defined]
        if index + 1 < len(lines):
            run.add_break()


def _set_list_level(target: object, level: int) -> None:
    if level == 0:
        return
    pPr = target._p.get_or_add_pPr()  # type: ignore[attr-defined]
    numPr = pPr.makeelement(qn("w:numPr"), {})
    ilvl = numPr.makeelement(qn("w:ilvl"), {qn("w:val"): str(level)})
    numId = numPr.makeelement(qn("w:numId"), {qn("w:val"): "1"})
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)
